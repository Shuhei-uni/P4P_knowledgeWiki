#!/usr/bin/env python3
"""Build a restrained, Google-Docs-style setup-07 supervisor report.

The document reuses only accepted figures that already exist in the setup-07
meeting evidence package. It does not run Fluent, alter a case, or create new
analysis graphics.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
FIGURE_DIR = ROOT / "output" / "setup07_meeting_visuals_20260811"
OUTPUT_DIR = ROOT / "output" / "docx"
WORKING = ROOT / "tmp" / "docx" / "setup07_supervisor_report_google_docs_style_working.docx"

FIGURES = {
    "baseline": FIGURE_DIR / "01_closed_bottom_inventory_pressure.png",
    "closed_liquid": FIGURE_DIR / "06_closed_bottom_liquid_field_comparison.png",
    "closed_pressure": FIGURE_DIR / "07_closed_bottom_pressure_field_comparison.png",
    "sink_swirl": FIGURE_DIR / "08_sink_band_and_carrier_swirl.png",
    "pathways": FIGURE_DIR / "21_setup07_liquid_pathway_accounting.png",
    "sink_liquid": FIGURE_DIR / "22_setup07c_setup07e_liquid_field_comparison.png",
    "residual": FIGURE_DIR / "15_setup07e_adaptive_residual_history.png",
    "outlet_plan": FIGURE_DIR / "25_brine_outlet_boundary_condition_decision.png",
    "outlet_sizing": FIGURE_DIR / "24_preliminary_brine_outlet_area_envelope.png",
}

BLACK = RGBColor(0x00, 0x00, 0x00)
GREY = RGBColor(0x55, 0x55, 0x55)
LIGHT_GREY = "DADCE0"


def set_font(run, size: float, *, bold: bool = False, italic: bool = False, color=BLACK) -> None:
    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def set_cell_margins(cell, *, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        tag = tc_mar.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            tc_mar.append(tag)
        tag.set(qn("w:w"), str(value))
        tag.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for name in ("top", "bottom", "insideH"):
        edge = OxmlElement(f"w:{name}")
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "6")
        edge.set(qn("w:space"), "0")
        edge.set(qn("w:color"), LIGHT_GREY)
        borders.append(edge)
    for name in ("start", "end", "insideV"):
        edge = OxmlElement(f"w:{name}")
        edge.set(qn("w:val"), "nil")
        borders.append(edge)


def set_table_geometry(table, widths_in: list[float]) -> None:
    widths_dxa = [round(width * 1440) for width in widths_in]
    total_dxa = sum(widths_dxa)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total_dxa))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "0")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width_dxa in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width_dxa))
        grid.append(grid_col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths_in[index])
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[index]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def mark_header_row(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = tr_pr.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        tr_pr.append(header)
    header.set(qn("w:val"), "true")


def set_picture_alt_text(inline_shape, description: str) -> None:
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("descr", description)
    doc_pr.set("title", description)


def add_body(doc: Document, text: str, *, bold_lead: str | None = None):
    paragraph = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        set_font(lead, 11, bold=True)
        rest = paragraph.add_run(text[len(bold_lead):])
        set_font(rest, 11)
    else:
        run = paragraph.add_run(text)
        set_font(run, 11)
    return paragraph


def add_bullet(doc: Document, text: str):
    paragraph = doc.add_paragraph(style="List Bullet")
    run = paragraph.add_run(text)
    set_font(run, 11)
    return paragraph


def add_numbered(doc: Document, text: str):
    paragraph = doc.add_paragraph(style="List Number")
    run = paragraph.add_run(text)
    set_font(run, 11)
    return paragraph


def add_equation(doc: Document, segments: list[tuple[str, bool]]) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(6)
    for text, subscript in segments:
        run = paragraph.add_run(text)
        set_font(run, 12, italic=True)
        run.font.subscript = subscript


def add_figure(doc: Document, path: Path, number: int, caption: str, source: str, *, width: float = 6.3) -> None:
    image_paragraph = doc.add_paragraph()
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.paragraph_format.keep_with_next = True
    inline_shape = image_paragraph.add_run().add_picture(str(path), width=Inches(width))
    set_picture_alt_text(inline_shape, caption)

    caption_paragraph = doc.add_paragraph()
    caption_paragraph.paragraph_format.space_before = Pt(3)
    caption_paragraph.paragraph_format.space_after = Pt(4)
    caption_paragraph.paragraph_format.keep_together = True
    caption_paragraph.paragraph_format.keep_with_next = False
    label = caption_paragraph.add_run(f"Figure {number}. ")
    set_font(label, 9, bold=True)
    text_run = caption_paragraph.add_run(caption)
    set_font(text_run, 9)
    source_run = caption_paragraph.add_run()
    source_run.add_break()
    source_run.add_text(f"Source: {source}")
    set_font(source_run, 8, color=GREY)


def format_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles

    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(11)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15

    heading_tokens = {
        "Heading 1": (20, BLACK, 20, 6),
        "Heading 2": (16, BLACK, 18, 6),
        "Heading 3": (14, RGBColor(0x43, 0x43, 0x43), 16, 4),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = False
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.15


def add_results_table(doc: Document) -> None:
    rows = [
        ("07a", "Closed bottom", "2.4%", "Carrier field remains unresolved; no physical liquid outlet."),
        ("07b", "One-cell sink", "8.8%", "Too little liquid removal; diagnostic only."),
        ("07c", "0.140 m fixed sink", "19.2%", "Removal improved, but most of the liquid route remained unclosed."),
        ("07d", "Stronger fixed sink", "40.2%", "Large increase in source strength produced diminishing return."),
        ("07e", "Adaptive sink", "43.6%", "Best tested routing, but zero convergence windows passed."),
    ]
    table = doc.add_table(rows=1, cols=4)
    headers = ("Branch", "Purpose", "Liquid routed", "Endpoint interpretation")
    for index, value in enumerate(headers):
        run = table.rows[0].cells[index].paragraphs[0].add_run(value)
        set_font(run, 9.5, bold=True)
    mark_header_row(table.rows[0])
    for row_values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row_values):
            paragraph = cells[index].paragraphs[0]
            run = paragraph.add_run(value)
            set_font(run, 9.5)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.08
            if index in (0, 2):
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_table_geometry(table, [0.65, 1.55, 1.05, 3.25])
    set_table_borders(table)


def add_terminal_table(doc: Document) -> None:
    rows = [
        ("Full-strength endpoint", "2,000 iterations", "Maximum diagnostic budget reached"),
        ("Numerical sink", "50.954 kg/s", "43.6% of the 116.92 kg/s liquid feed"),
        ("Liquid imbalance", "56.42%", "Complete phase balance not achieved"),
        ("Pressure drop", "26.725 kPa", "5.85% drift over the final 500 iterations"),
        ("Liquid inventory", "65.008 kg", "13.35% drift over the final 500 iterations"),
        ("Continuity residual", "0.254617", "Failed the 1e-3 gate and rose near the endpoint"),
        ("Liquid VF residual", "7.53e-4", "Below 1e-3, but rising near the endpoint"),
        ("Accepted windows", "0", "Iteration independence was not established"),
    ]
    table = doc.add_table(rows=1, cols=3)
    headers = ("Metric", "Terminal value", "Interpretation")
    for index, value in enumerate(headers):
        run = table.rows[0].cells[index].paragraphs[0].add_run(value)
        set_font(run, 9.5, bold=True)
    mark_header_row(table.rows[0])
    for row_values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row_values):
            paragraph = cells[index].paragraphs[0]
            run = paragraph.add_run(value)
            set_font(run, 9.5)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.08
            if index == 1:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_table_geometry(table, [1.65, 1.4, 3.45])
    set_table_borders(table)


def build() -> Path:
    missing = [str(path) for path in FIGURES.values() if not path.is_file()]
    if missing:
        raise FileNotFoundError(f"Missing accepted figure(s): {missing}")

    WORKING.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    doc = Document()
    format_document(doc)
    doc.core_properties.title = "Geothermal Steam-Water Separator CFD: Setup 07 Diagnostic Findings"
    doc.core_properties.subject = "Condensed supervisor report on the carrier-field and liquid-outlet diagnostics"
    doc.core_properties.author = "P4P geothermal separator CFD project"
    doc.core_properties.keywords = "CFD, Fluent, geothermal separator, brine outlet, mesh convergence"

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(3)
    title_run = title.add_run("Geothermal Separator CFD")
    set_font(title_run, 26, bold=False)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    subtitle_run = subtitle.add_run("Setup 07 diagnostic findings and recommended brine-outlet rebuild")
    set_font(subtitle_run, 14, color=GREY)

    metadata = doc.add_paragraph()
    metadata.paragraph_format.space_after = Pt(18)
    metadata_run = metadata.add_run("Supervisor meeting report | 11 August 2026 | Carrier model only; DPM and EWF off")
    set_font(metadata_run, 10, color=GREY)

    doc.add_heading("Executive summary", level=1)
    add_body(
        doc,
        "Recommendation. Add a resolved brine pipe and qualify one freshly initialized 900k carrier case before repeating mesh convergence or adding DPM/EWF.",
        bold_lead="Recommendation.",
    )
    add_body(
        doc,
        "The present split-inlet model has a steam pressure outlet, while the surface representing the water level is a wall. The liquid phase therefore has no physical brine discharge. The steady solver continues to redistribute and retain liquid, so the apparent solution changes with iteration even after thousands of iterations.",
    )
    add_body(
        doc,
        "Four local sink strategies were tested as numerical surrogates for a constant water level. They increased liquid removal, but none closed the complete liquid route or passed the agreed residual and physical-monitor stability criteria. The best case, setup 07e, removed 50.95 kg/s (43.6% of the liquid feed) and still ended with a 56.42% liquid imbalance and zero accepted convergence windows.",
    )

    doc.add_heading("What the study does establish", level=2)
    for text in (
        "Seven carrier-field mesh runs were completed, but continuing within-run drift prevents a defensible mesh-independence claim.",
        "The closed-bottom geometry is the dominant limitation: liquid inventory and pressure continue to change because the model does not include outlet hydraulics for the brine.",
        "Local sink sources are useful diagnostics, but their result is controlled by liquid transport into the sink band and cannot be interpreted as a validated drain boundary.",
        "Steam quality and liquid carryover from these cases remain trend-only quantities rather than validated separator-efficiency predictions.",
    ):
        add_bullet(doc, text)

    pace = doc.add_paragraph()
    pace.paragraph_format.space_before = Pt(8)
    pace_run = pace.add_run("Suggested 10-minute walkthrough: ")
    set_font(pace_run, 10, bold=True)
    pace_detail = pace.add_run("1 min summary; 2 min model limitation; 2 min Fluent fields; 2 min sink tests; 1 min convergence; 2 min recommendation.")
    set_font(pace_detail, 10)

    doc.add_page_break()
    doc.add_heading("1. Model basis and diagnostic branches", level=1)
    add_body(
        doc,
        "All comparisons use the same post-replication split-inlet separator geometry and carrier physics. The model is pressure-based and steady, with the Mixture multiphase formulation, vapor as the primary phase, liquid water as the secondary phase, RNG k-epsilon turbulence, gravity (0, -9.81, 0), Energy off and SIMPLE/PRESTO coupling. The liquid and steam feeds are 116.92 and 80.69 kg/s, and the steam outlet pressure is 1.12 MPa. The bottom surface remains a wall. DPM and EWF were not run.",
    )
    add_body(
        doc,
        "The 900k mesh was used for the extended diagnostic branches because it was computationally practical and already showed the underlying model problem. These branches were designed to test whether a numerical sink could emulate Purnanto's assumed constant water level without changing the geometry.",
    )
    doc.add_heading("Implemented numerical sink", level=2)
    add_body(
        doc,
        "In the marked cells, the UDF removes only secondary-phase liquid mass using the following volumetric source:",
    )
    add_equation(
        doc,
        [
            ("S", False), ("l", True), (" = -rho", False), ("l", True),
            (" alpha", False), ("l", True), (" R / tau     [kg m^-3 s^-1]", False),
        ],
    )
    add_body(
        doc,
        "Here, rho_l is the liquid density, alpha_l is the local liquid volume fraction, R is the guarded ramp factor from 0 to 1, and tau is the removal time scale. The negative sign denotes removal. Integrating the source over the active band gives:",
    )
    add_equation(
        doc,
        [
            ("m_dot", False), ("sink", True), (" = -(R / tau) integral over ", False),
            ("band", True), ("(rho", False), ("l", True), (" alpha", False),
            ("l", True), (" dV) = -R M", False), ("l,band", True), (" / tau", False),
        ],
    )
    add_body(
        doc,
        "The UDF also removes the mixture momentum carried by the deleted liquid through S_i = S_l u_m,i. It applies no vapor, energy or DPM source. Setup 07e updated the removal time scale every 100 iterations using the bounded feedback law below.",
    )
    add_equation(
        doc,
        [
            ("tau", False), ("next", True), (" = clamp(M", False), ("l,band", True),
            (" / 116.92, 0.002, 0.2) s", False),
        ],
    )
    add_body(
        doc,
        "At the minimum tau, the active band contained 0.101909 kg, so the law could remove about 50.95 kg/s; removing the full liquid feed would have required 0.23384 kg in the band. This is why the adaptive sink became locally capacity-limited.",
    )
    doc.add_page_break()
    doc.add_heading("Diagnostic sequence", level=2)
    add_results_table(doc)
    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(8)
    note_run = note.add_run("Interpretation note. ")
    set_font(note_run, 10, bold=True)
    note_text = note.add_run("\"Liquid routed\" combines liquid leaving through the steam outlet and liquid removed by the numerical source. It is an endpoint mass-accounting diagnostic, not proof of a physical drain or transient level response.")
    set_font(note_text, 10)

    doc.add_heading("Study acceptance logic", level=2)
    add_body(
        doc,
        "Residuals were assessed together with pressure drop, liquid inventory, outlet phase flow, velocity, swirl/vorticity and complete phase/mixture balance. A fixed iteration count was not treated as convergence. Mesh independence was to be assessed only after an iteration-independent carrier solution was established.",
    )
    doc.add_heading("Questions to resolve in the meeting", level=2)
    for text in (
        "Confirm the physical drain location, target water level and available downstream brine pressure.",
        "Confirm a preliminary outlet diameter. The continuity-only starting envelope is 0.411, 0.291 or 0.237 m at mean liquid velocities of 1, 2 or 3 m/s.",
        "Approve a pressure-outlet-first qualification and the rule that mesh convergence, DPM and EWF remain on hold until carrier mass closure and stability pass.",
    ):
        add_bullet(doc, text)

    doc.add_page_break()
    doc.add_heading("2. Closed-bottom baseline: continuing accumulation and redistribution", level=1)
    add_body(
        doc,
        "The clearest evidence is the extended 900k closed-bottom run. Between iterations 4000 and 6000, the liquid inventory increased from 104.05 to 171.03 kg (+64.4%), while the pressure drop increased from 31.03 to 34.05 kPa (+9.7%). The change is much larger than an acceptable final-window drift and is comparable to, or larger than, several mesh-to-mesh differences.",
    )
    add_figure(
        doc,
        FIGURES["baseline"],
        1,
        "Liquid inventory and pressure drop continue to change during the closed-bottom 900k extension.",
        "Saved 900k Fluent case/data checkpoints and terminal monitor reports at iterations 4000-6000.",
        width=6.3,
    )
    add_body(
        doc,
        "Because the underlying field is still moving with solver iteration, Richardson extrapolation and Grid Convergence Index calculations are not defensible for this sequence. More iterations alone also cannot repair the missing physical liquid route: the steady solution is trying to accommodate a continuous liquid feed in a domain with no brine outlet.",
    )

    doc.add_page_break()
    doc.add_heading("3. What the Fluent fields show", level=1)
    add_body(
        doc,
        "Matched Fluent sections at x = -1.5 m show where the changing integral quantities are expressed spatially. The liquid-volume-fraction plots use the same clipped 0-0.05 range to reveal low-volume liquid structure. They do not represent a free surface, and the iteration axis is not physical time.",
    )
    add_figure(
        doc,
        FIGURES["closed_liquid"],
        2,
        "The liquid-rich wall region expands between iterations 4000 and 6000 as the domain inventory rises.",
        "Fluent phase-2 liquid-volume-fraction contours on the same 900k section and fixed display range.",
        width=6.2,
    )
    add_body(
        doc,
        "Matched absolute-pressure contours also shift over the same interval, consistent with the monitored pressure-drop increase from 31.03 to 34.05 kPa. The liquid field is the more direct visual evidence for the present geometry problem: the wall-adjacent liquid-rich region grows because the model has no resolved brine discharge.",
    )

    doc.add_page_break()
    doc.add_heading("4. Local sink tests improve removal but do not close the liquid route", level=1)
    add_body(
        doc,
        "The sink sequence increased the represented liquid route from 8.8% for the one-cell sink to 43.6% for the adaptive thick-band sink. However, even the strongest case left approximately 66 kg/s of the steady liquid route unclosed. The improvement also showed diminishing returns as the active band became inventory-limited.",
    )
    add_figure(
        doc,
        FIGURES["pathways"],
        3,
        "Endpoint liquid-pathway accounting for the closed-bottom and sink branches.",
        "Corrected source-inclusive phase mass reports. The grey portion is an unclosed steady rate, not a measured physical-time accumulation rate.",
        width=6.3,
    )
    add_body(
        doc,
        "This is the main reason not to keep increasing the numerical source coefficient. The source can remove only liquid that the carrier field transports into the selected band. Making the source stronger dries that limited region faster; it does not create the pressure-driven, area-resolved discharge that a brine pipe would provide.",
    )

    doc.add_page_break()
    doc.add_heading("5. Sink placement and carrier-flow evidence", level=1)
    add_body(
        doc,
        "The thick sink acts only in a 0.140165 m band adjacent to the bottom wall. The Mixture-carrier pathlines released from the liquid inlet show strong swirl and recirculation through the vessel. These are carrier trajectories coloured by velocity; they are not droplets and were generated with DPM off.",
    )
    add_figure(
        doc,
        FIGURES["sink_swirl"],
        4,
        "Fluent sink-mask and liquid-inlet carrier pathlines for setup 07c.",
        "Saved setup-07c endpoint; fixed bottom-local source band and Mixture-carrier pathlines.",
        width=6.2,
    )
    add_body(
        doc,
        "Matched phase-2 contours from the fixed and adaptive endpoints were also inspected on the same plane and display range. The local fields differ, but neither case forms a resolved outlet flow. This supports treating both source variants as numerical diagnostics rather than alternative physical drain models.",
    )

    doc.add_page_break()
    doc.add_heading("6. Convergence assessment of the best sink case", level=1)
    add_body(
        doc,
        "Setup 07e was the most capable surrogate because it adaptively strengthened the sink down to the implemented minimum time scale. It completed the full iteration budget, but the continuity residual remained high and the physical monitors were still moving. The liquid-volume-fraction residual falling below 1e-3 was therefore not sufficient to accept the solution.",
    )
    add_figure(
        doc,
        FIGURES["residual"],
        5,
        "Residual response of the completed setup-07e adaptive sink diagnostic.",
        "Complete residual history from the guarded 1000-iteration ramp and 2000 full-strength iterations.",
        width=6.15,
    )
    add_terminal_table(doc)

    doc.add_page_break()
    doc.add_heading("7. Recommended rebuild and next qualification", level=1)
    add_body(
        doc,
        "The next model should represent the brine discharge explicitly. The outlet face should be placed at the end of a resolved pipe or extension, away from the vessel recirculation zone. If the downstream static pressure is known or defensibly estimated, a pressure outlet is the preferred first case because Fluent can solve the discharged mass flow from the field.",
    )
    add_figure(
        doc,
        FIGURES["outlet_sizing"],
        6,
        "Preliminary continuity-only outlet-size envelope for the 116.92 kg/s liquid feed.",
        "Existing project planning figure using the Fluent-read liquid density of 881.211 kg/m3. It excludes pressure losses, flashing/cavitation, back-pressure, level control and multiphase outlet behaviour.",
        width=6.05,
    )

    doc.add_heading("Boundary-condition approach", level=2)
    add_body(
        doc,
        "Use a pressure outlet first when downstream static pressure is known or defensibly estimated. A 116.92 kg/s mass-flow outlet can be retained as a strictly outward diagnostic bracket where the required rate is known but the downstream pressure is not. Avoid an outflow boundary for the first rebuild because the vessel already shows recirculation and outlet-reversal sensitivity.",
    )

    doc.add_heading("Proposed qualification sequence", level=2)
    for text in (
        "Create a resolved brine pipe and named brine-outlet face away from the vessel recirculation region.",
        "Apply the pressure-outlet-first approach above, specify phase backflow consistently and monitor any reversal.",
        "Freshly initialize one 900k carrier case with the accepted carrier physics and require complete mass closure plus two stable monitor windows.",
        "Repeat the mesh sequence only after that carrier case passes; keep DPM and EWF off until then.",
    ):
        add_numbered(doc, text)

    doc.save(WORKING)
    return WORKING


if __name__ == "__main__":
    print(build())
