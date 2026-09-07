from copy import deepcopy

from docx import Document
from docx.shared import Pt


DOCX_PATH = "/Users/andy/Desktop/P4P/Meeting Minutes - P4P.docx"


def copy_paragraph_properties(source, target):
    source_ppr = source._p.pPr
    if source_ppr is None:
        return
    target_ppr = target._p.get_or_add_pPr()
    target_ppr.clear()
    for child in source_ppr:
        target_ppr.append(deepcopy(child))


def remove_paragraph(paragraph):
    paragraph._element.getparent().remove(paragraph._element)


def add_section_heading(doc, text):
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt(18)
    return paragraph


def add_body(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.add_run(text)
    return paragraph


def add_bullet(doc, text, bullet_template):
    paragraph = doc.add_paragraph()
    copy_paragraph_properties(bullet_template, paragraph)
    paragraph.add_run(text)
    return paragraph


doc = Document(DOCX_PATH)

wk12_index = None
for index, paragraph in enumerate(doc.paragraphs):
    if paragraph.text == "WK 12 - Droplet Injection Simulation Review":
        wk12_index = index
        break

if wk12_index is None:
    raise RuntimeError("Could not find WK 12 heading")

next_week_index = None
for index in range(wk12_index + 1, len(doc.paragraphs)):
    text = doc.paragraphs[index].text
    if text.startswith("WK ") and index != wk12_index:
        next_week_index = index
        break

section_end = next_week_index if next_week_index is not None else len(doc.paragraphs)
for paragraph in list(doc.paragraphs[wk12_index + 1 : section_end]):
    remove_paragraph(paragraph)

bullet_template = None
for paragraph in doc.paragraphs:
    ppr = paragraph._p.pPr
    if ppr is not None and ppr.numPr is not None:
        bullet_template = paragraph
        break

if bullet_template is None:
    raise RuntimeError("Could not find a bullet template")

add_section_heading(doc, "Summary")
add_body(
    doc,
    "The team reviewed droplet-injection results from the professional setup. "
    "The flux result suggests very low liquid carryover and about 99.95% efficiency, "
    "but the run still needs better mass balance and convergence evidence before it "
    "can support report-level claims. The DPM results showed a non-monotonic droplet-size "
    "trend, especially the higher escape rate for 5 micrometer droplets compared with "
    "1 micrometer and 10 micrometer droplets.",
)

add_section_heading(doc, "Insights")
for insight in [
    "Diagnostic-only: this is useful steam-line carryover evidence, not full separator validation yet.",
    "Reported: flux efficiency was about 99.95%, but the meeting flagged imbalance between inlet and outlet steam flow.",
    "Inferred: raw particle-count escape should stay a path diagnostic; standard efficiency, outlet dryness, and mass-weighted carryover are the defensible comparison metrics.",
    "Inferred: the 5 micrometer escape peak is suspicious because smaller droplets should generally follow the steam more easily; check mesh, stochastic tracking, rotation, drag law, and sample size.",
    "Inferred: wall/top-ring accumulation and possible re-entrainment may be important, so wall fate assumptions need to be stated clearly.",
]:
    add_bullet(doc, insight, bullet_template)

add_section_heading(doc, "Action Items")
for action in [
    "Improve carrier-flow balance and convergence before using efficiency values in the report.",
    "Report standard separator efficiency and outlet dryness alongside the count-based DPM escape result.",
    "Expand the DPM sweep across 10-50 micrometer droplets, including 20 micrometer and 50 micrometer cases.",
    "Test DPM settings sensitivity: discrete random walk, rotation, drag law, and particle density.",
    "Andy - get help locating the referenced GRC Transactions paper and package the Workbench project as a single archive file.",
]:
    add_bullet(doc, action, bullet_template)

doc.save(DOCX_PATH)
