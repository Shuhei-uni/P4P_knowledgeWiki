from copy import deepcopy

from docx import Document
from docx.shared import Pt


DOCX_PATH = "/Users/andy/Desktop/P4P/Meeting Minutes - P4P.docx"


def set_paragraph_text(paragraph, text, size_pt=None, bold=None):
    paragraph.clear()
    run = paragraph.add_run(text)
    if size_pt is not None:
        run.font.size = Pt(size_pt)
    if bold is not None:
        run.bold = bold


def copy_paragraph_properties(source, target):
    source_ppr = source._p.pPr
    if source_ppr is None:
        return
    target_ppr = target._p.get_or_add_pPr()
    target_ppr.clear()
    for child in source_ppr:
        target_ppr.append(deepcopy(child))


def add_heading(doc, text, heading_template):
    paragraph = doc.add_paragraph()
    copy_paragraph_properties(heading_template, paragraph)
    set_paragraph_text(paragraph, text, size_pt=14)
    return paragraph


def add_body(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.add_run(text)
    return paragraph


def add_action_heading(doc):
    return add_section_heading(doc, "Action Items")


def add_section_heading(doc, text):
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt(18)
    return paragraph


def add_insight_heading(doc):
    return add_section_heading(doc, "Insights")


def add_action_item(doc, text, bullet_template):
    paragraph = doc.add_paragraph()
    copy_paragraph_properties(bullet_template, paragraph)
    paragraph.add_run(text)
    return paragraph


def add_entry(doc, heading, summary, actions, heading_template, bullet_template, insights=None):
    doc.add_paragraph("")
    add_heading(doc, heading, heading_template)
    if isinstance(summary, dict):
        for section_title, section_value in summary.items():
            add_section_heading(doc, section_title)
            if isinstance(section_value, list):
                for item in section_value:
                    add_action_item(doc, item, bullet_template)
            else:
                add_body(doc, section_value)
        if actions:
            add_action_heading(doc)
            for action in actions:
                add_action_item(doc, action, bullet_template)
        return

    add_body(doc, summary)
    if insights:
        add_insight_heading(doc)
        for insight in insights:
            add_action_item(doc, insight, bullet_template)
    if actions:
        add_action_heading(doc)
        for action in actions:
            add_action_item(doc, action, bullet_template)


doc = Document(DOCX_PATH)

# Normalize the existing later headings to the requested "WK # - Theme" pattern.
heading_updates = {
    "Week 5 – 02/04": "WK 5 - Two-Phase Flow Modeling Meeting",
    "Week 5 – 02/04 ": "WK 5 - Two-Phase Flow Modeling Meeting",
    "Week 6 – 23/04": "WK 6 - Spiral Inlet Design Meeting",
    "Week 7 – 30/04": "WK 7 - Steam-Liquid Separator Analysis",
}

for paragraph in doc.paragraphs:
    current = paragraph.text
    if current in heading_updates:
        set_paragraph_text(paragraph, heading_updates[current], size_pt=14)

start_remove = None
for index, paragraph in enumerate(doc.paragraphs):
    if paragraph.text.startswith("WK 8 -"):
        start_remove = index
        break

if start_remove is not None:
    for paragraph in list(doc.paragraphs[start_remove:]):
        paragraph._element.getparent().remove(paragraph._element)

while doc.paragraphs and not doc.paragraphs[-1].text.strip():
    paragraph = doc.paragraphs[-1]
    paragraph._element.getparent().remove(paragraph._element)

heading_template = next(
    paragraph
    for paragraph in doc.paragraphs
    if paragraph.text == "WK 7 - Steam-Liquid Separator Analysis"
)
bullet_template = next(
    paragraph
    for paragraph in doc.paragraphs
    if paragraph.text.startswith("Email the specific file name")
)

entries = [
    {
        "heading": "WK 8 - CFD Model Troubleshooting and Data Access",
        "summary": (
            "The meeting focused on issues with a CFD model involving water and steam. "
            "The team discussed initializing the water level, turbulent mixing, and the "
            "difficulty of comparing against the previous student's model because access "
            "to the Google Drive files had not been granted. The inlet boundary condition "
            "was adjusted from mass flow to velocity inlet because the steam and water "
            "velocities differed strongly. The team also discussed using the university "
            "Copilot account for technical support and using field data or the empirical "
            "1984 paper for later validation."
        ),
        "actions": [
            "If upcoming tests prevent project progress, email the supervisor to cancel or reschedule the next meeting.",
            "Use the university-provided Copilot account instead of a personal ChatGPT account for coding and technical assistance.",
            "Continue following up with the former student to obtain access to the original model and data files.",
            "After access is granted, compare the current CFD simulation results with the student's model to identify discrepancies.",
            "Follow up with Sally to obtain experimental field data from the company.",
            "After gaining confidence in the CFD model, compare simulation results with company field data or the empirical 1984 paper.",
        ],
    },
    {
        "heading": "WK 9 - Missed Meeting",
        "summary": (
            "No formal meeting minutes were recorded for this week because the meeting was missed."
        ),
        "actions": [],
    },
    {
        "heading": "WK 10 - Two-Phase Inlet and Steam Quality Modeling",
        "summary": (
            "The team discussed challenges in modeling a two-phase inlet system using "
            "separate water and steam streams. The setup used velocity inlets and a "
            "fuller geometry including the brine outlet, but convergence errors, water "
            "pooling, and outlet turbulence remained issues. The meeting clarified that "
            "the main research focus should stay on the inlet condition and its effect "
            "on steam quality and carryover, rather than overcomplicating the full water "
            "pool model. The team also reviewed mass flow, enthalpy, dryness fraction, "
            "separator efficiency, and the need to track small droplets in the steam."
        ),
        "actions": [
            "Update the Fluent inlet boundary condition so steam and water enter as separate phases, with water on the outside and steam on the inside.",
            "Reduce the pressure at the brine outlet to prevent flooding and stabilize the bottom water level.",
            "Bring contour plots and other visualizations from the current Fluent simulations to the next meeting.",
            "Ask Maida about appropriate Fluent solver settings, including energy, mixture model, and turbulence model choices.",
            "Measure the steam outlet geometry from the research-paper figure using a ruler to obtain approximate dimensions.",
            "If the water-pool initialization cannot be stabilized, revert to the earlier simplified initialization approach.",
            "Share the Fluent case files for the 2566-iteration and 4000-iteration simulations via Google Drive.",
        ],
    },
    {
        "heading": "WK 11 - Model Validation and DPM Setup Review",
        "summary": (
            "The team reviewed problems with the DPM particle-injection setup. Although "
            "the flux report suggested approximately 97% separator efficiency and some "
            "liquid carryover, injected particles were not escaping through the steam "
            "outlet as expected from the referenced model. The group discussed correcting "
            "the inlet setup so pure steam and pure liquid enter through area ratios "
            "based on volumetric flow rates, comparing flux-based efficiency with the "
            "analytical model, and asking Maida to review the particle-injection setup."
        ),
        "actions": [
            "Ask Maida for suggestions on why the current CFD model is not behaving like the Purnanto model, focusing on particle injection and the overall setup.",
            "Set up an Excel sheet for the analytical efficiency model from the paper so simulation parameters can be compared against it.",
            "Update the CFD inlet setup to use separate pure steam and pure liquid inlets with area ratios proportional to volumetric flow rates.",
            "Coordinate with Maida to review the CFD model and particle-injection setup, including sharing the model for diagnosis.",
            "Run a trial CFD simulation injecting small droplets directly into the steam inlet to see whether any escape through the steam outlet.",
            "Compute separator efficiency from the flux report and compare it against the analytical efficiency model.",
        ],
    },
    {
        "heading": "WK 12 - Droplet Injection Simulation Review",
        "summary": {
            "Summary": (
                "The team reviewed the latest droplet-injection simulations for the "
                "professional setup. The flux result suggests very low liquid carryover, "
                "but the run still needs mass-balance and convergence checks before it "
                "can support report-level efficiency claims. The DPM results showed a "
                "non-monotonic droplet-size trend, with 5 micrometer droplets escaping "
                "more often than 1 micrometer droplets."
            ),
            "Decisions / Interpretation": [
                "Diagnostic-only: treat this setup as steam-line carryover evidence, not full separator validation.",
                "Reported: flux-based efficiency was around 99.95%, but mass balance was still a concern.",
                "Inferred: particle-count escape is a path diagnostic; use mass-weighted carryover or outlet quality for Purnanto-style comparison.",
                "Inferred: the 5 micrometer escape peak is suspicious and should be checked with controlled DPM, mesh, and turbulence-setting sensitivities.",
            ],
            "Action Items": [
                "CFD model owner - include droplet injection with the steam inlet and rerun the DPM checks.",
                "CFD model owner - improve mesh/solver balance before using efficiency values in the report.",
                "CFD model owner - report standard separator efficiency alongside particle-count escape results.",
                "Andy - archive the Workbench project and bring/send the single archive file for review.",
            ],
            "Next Focus": [
                "Run 5, 10, and 40-50 micrometer DPM cases with the same carrier-flow field.",
                "Record escaped, trapped, and incomplete counts plus the represented mass basis.",
                "Check whether slower proportional inlet flow increases residence time and improves separation.",
            ],
        },
        "actions": [
        ],
    },
]

for entry in entries:
    add_entry(
        doc,
        entry["heading"],
        entry["summary"],
        entry["actions"],
        heading_template,
        bullet_template,
        entry.get("insights"),
    )

doc.save(DOCX_PATH)
